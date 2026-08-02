#!/usr/bin/env python3
"""드래곤빌리지3 엑셀 원본 → 사이트용 JSON 데이터 생성 스크립트.

scripts/dragons.xlsx, scripts/orbs.xlsx 를 읽어 data/*.json 으로 출력한다.
원본 갱신 시 `python3 scripts/build_data.py` 를 다시 실행하면 된다.
"""
import json
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
os.makedirs(DATA, exist_ok=True)


def rows(ws):
    return [r for r in ws.iter_rows(values_only=True)]


def sheet_records(ws):
    data = rows(ws)
    header = [str(h).strip() if h is not None else "" for h in data[0]]
    out = []
    for r in data[1:]:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        rec = {}
        for k, v in zip(header, r):
            rec[k] = v
        out.append(rec)
    return out


def dump(name, obj):
    path = os.path.join(DATA, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    print(f"wrote {path} ({len(obj) if isinstance(obj, list) else 'obj'})")


def build_dragons():
    wb = openpyxl.load_workbook(os.path.join(HERE, "dragons.xlsx"), data_only=True)
    dragons = []
    for rec in sheet_records(wb["드래곤 목록"]):
        dragons.append({
            "element": rec["주속성"],
            "subElements": [s.strip() for s in str(rec["부속성"]).split("·")] if rec.get("부속성") else [],
            "grade": rec["등급"],
            "name": rec["이름"],
            "attackType": rec["공격타입"],
            "total": rec["Lv.50 합계"],
            "stats": {
                "hp": rec["체력"], "atk": rec["공격"], "def": rec["방어"],
                "mag": rec["마력"], "res": rec["저항"], "spd": rec["속도"],
            },
            "ability": rec["특성(어빌리티)"],
            "abilityEffect": rec["특성 효과"],
            "basicAttack": rec["평타"],
            "ultimate": rec["필살기"],
            "ultimateType": rec["필살기타입"],
            "accuracy": rec["명중"],
            "power": rec["위력"],
            "ultimateEffect": rec["필살기 효과"],
            "source": rec["획득처"],
        })

    abilities = []
    for rec in sheet_records(wb["특성(어빌리티) 사전"]):
        abilities.append({
            "name": rec["특성명"],
            "effect": rec["효과"],
            "count": rec["보유 수"],
            "dragons": [d.strip() for d in str(rec["보유 드래곤"]).split("/")],
        })

    summary = sheet_records(wb["속성·등급별 요약"])
    notes = [r[0] for r in rows(wb["안내"]) if r and r[0]]

    dump("dragons.json", dragons)
    dump("abilities.json", abilities)
    dump("dragon_summary.json", summary)
    dump("dragon_notes.json", notes)


def build_orbs():
    wb = openpyxl.load_workbook(os.path.join(HERE, "orbs.xlsx"), data_only=True)
    orbs = []
    for rec in sheet_records(wb["보주 목록"]):
        orbs.append({
            "element": rec["속성"],
            "grade": rec["등급"],
            "name": rec["보주명"],
            "skill": rec["스킬명"],
            "type": rec["타입"],
            "accuracy": rec["명중"],
            "power": rec["위력"],
            "cost": rec["비용(게이지)"],
            "statBonus": rec["핵심능력치 보너스"],
            "effect": rec["효과"],
            "source": rec["주요 획득처"],
        })
    summary = sheet_records(wb["속성별 요약"])
    terms = sheet_records(wb["상태이상·날씨 용어"])
    notes = [r[0] for r in rows(wb["안내"]) if r and r[0]]

    dump("orbs.json", orbs)
    dump("orb_summary.json", summary)
    dump("terms.json", terms)
    dump("orb_notes.json", notes)


if __name__ == "__main__":
    build_dragons()
    build_orbs()
    print("done.")
